"""Summary: produce MailArchivingSummary.xlsx and .txt."""
from __future__ import annotations

import logging
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from docflow.classification.classifier import FilingDecision

logger = logging.getLogger(__name__)

HEADERS = [
    "File Name",
    "Folder Location",
    "Document Type",
    "Institution",
    "Time Period",
    "Pages",
    "Rule Matched",
    "Confidence",
]


def _decision_to_row(decision: FilingDecision) -> list:
    """Convert a FilingDecision to a spreadsheet row."""
    return [
        decision.filename,
        decision.target_directory,
        decision.candidate.doc_type or "Unknown",
        decision.candidate.institution or "Unknown",
        decision.candidate.period or "Unknown",
        len(decision.candidate.pages),
        decision.rule_matched,
        round(decision.confidence, 2),
    ]


def generate_summary(
    auto_filed: list[FilingDecision],
    review_queue: list[FilingDecision],
    config: dict,
) -> tuple[Path, Path]:
    """Write summary spreadsheet and text file to archive root.

    Returns (xlsx_path, txt_path).
    """
    archive_root = Path(os.path.expanduser(
        config.get("archive_root", "~/DocFlowExample/archive")
    ))
    archive_root.mkdir(parents=True, exist_ok=True)

    xlsx_path = archive_root / "MailArchivingSummary.xlsx"
    txt_path = archive_root / "MailArchivingSummary.txt"

    all_decisions = auto_filed + review_queue
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # --- XLSX ---
    if xlsx_path.exists():
        wb = load_workbook(xlsx_path)
        ws = wb.active
        # Unmerge any merged cells from previous versions to avoid
        # MergedCell errors when iterating columns
        for merge_range in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merge_range))
        # Ensure headers exist (in case of a corrupted file)
        if ws.max_row == 0 or ws.cell(1, 1).value != HEADERS[0]:
            ws.insert_rows(1)
            for i, h in enumerate(HEADERS, 1):
                ws.cell(1, i, h)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Filing Summary"
        ws.append(HEADERS)

    for decision in all_decisions:
        row = _decision_to_row(decision)
        # Mark review queue items
        if not decision.auto_file:
            row[0] = f"[REVIEW] {row[0]}"
        ws.append(row)

    # Auto-size columns (approximate)
    try:
        for col in ws.columns:
            cells = list(col)
            if not cells:
                continue
            max_len = max(len(str(cell.value or "")) for cell in cells)
            letter = cells[0].column_letter
            ws.column_dimensions[letter].width = min(max_len + 2, 50)
    except Exception:
        # Column sizing is cosmetic — don't let it crash the pipeline
        logger.warning("Could not auto-size columns in summary spreadsheet")

    wb.save(xlsx_path)
    logger.info("Summary XLSX written: %s", xlsx_path)

    # --- TXT ---
    mode = "a" if txt_path.exists() else "w"
    with open(txt_path, mode) as f:
        f.write(f"\n{'=' * 60}\n")
        f.write(f"Mail Archiving Summary — {timestamp}\n")
        f.write(f"{'=' * 60}\n\n")

        f.write(f"Auto-filed: {len(auto_filed)}\n")
        f.write(f"Review queue: {len(review_queue)}\n")
        f.write(f"Total: {len(all_decisions)}\n\n")

        for decision in auto_filed:
            f.write(
                f"  [FILED] {decision.filename}\n"
                f"          → {decision.target_directory}\n"
                f"          Rule: {decision.rule_matched} "
                f"(confidence: {decision.confidence:.2f})\n\n"
            )

        for decision in review_queue:
            f.write(
                f"  [REVIEW] {decision.filename}\n"
                f"           → {decision.target_directory}\n"
                f"           Rule: {decision.rule_matched} "
                f"(confidence: {decision.confidence:.2f})\n"
                f"           Notes: {decision.notes or 'none'}\n\n"
            )

    logger.info("Summary TXT written: %s", txt_path)

    return xlsx_path, txt_path
