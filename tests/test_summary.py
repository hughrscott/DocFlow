"""Tests for src/summary/generator."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from docflow.clustering.clusterer import DocumentCandidate
from docflow.classification.classifier import FilingDecision
from docflow.summary.generator import generate_summary


def _make_decision(auto_file: bool = True) -> FilingDecision:
    candidate = DocumentCandidate(
        pages=[1, 2], institution="pnc", account="7364",
        period="February2026", doc_type="statement",
        clustering_confidence=0.9,
    )
    return FilingDecision(
        candidate=candidate,
        filename="PNCStatement.pdf",
        target_directory="/tmp/test",
        rule_matched="pnc_rule",
        confidence=0.9 if auto_file else 0.4,
        auto_file=auto_file,
        notes=None if auto_file else "Low confidence",
    )


class TestGenerateSummary:
    def test_creates_xlsx_and_txt(self, tmp_path):
        config = {"archive_root": str(tmp_path)}
        auto = [_make_decision(True)]
        review = [_make_decision(False)]
        xlsx, txt = generate_summary(auto, review, config)

        assert xlsx.exists()
        assert txt.exists()
        assert xlsx.suffix == ".xlsx"
        assert txt.suffix == ".txt"

    def test_xlsx_has_correct_rows(self, tmp_path):
        config = {"archive_root": str(tmp_path)}
        auto = [_make_decision(True), _make_decision(True)]
        review = [_make_decision(False)]
        xlsx, _ = generate_summary(auto, review, config)

        wb = load_workbook(xlsx)
        ws = wb.active
        # 1 header + 3 data rows
        assert ws.max_row == 4

    def test_txt_contains_summary(self, tmp_path):
        config = {"archive_root": str(tmp_path)}
        auto = [_make_decision(True)]
        review = []
        _, txt = generate_summary(auto, review, config)

        content = txt.read_text()
        assert "Auto-filed: 1" in content
        assert "PNCStatement.pdf" in content

    def test_appends_to_existing(self, tmp_path):
        config = {"archive_root": str(tmp_path)}
        auto = [_make_decision(True)]

        generate_summary(auto, [], config)
        generate_summary(auto, [], config)

        xlsx = tmp_path / "MailArchivingSummary.xlsx"
        wb = load_workbook(xlsx)
        ws = wb.active
        # 1 header + 2 runs of 1 row each
        assert ws.max_row == 3
